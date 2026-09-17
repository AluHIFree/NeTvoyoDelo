"""
Красивый экспорт корреспонденции в Excel (openpyxl).
Цвета статусов и сроков, шапка, автофильтр, чередование строк.
"""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from typing import Any, Iterable, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app import models
from app.services.productivity import deadline_stage


# --- Палитра (без «кричащего» фиолетового) ---
HEADER_FILL = PatternFill("solid", fgColor="1A4C8C")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1A4C8C")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="64748B")
CELL_FONT = Font(name="Calibri", size=10, color="1E293B")
ZEBRA_FILL = PatternFill("solid", fgColor="F8FAFC")
THIN = Side(style="thin", color="CBD5E1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

STATUS_FILLS = {
    "В ожидании": PatternFill("solid", fgColor="FEF3C7"),
    "В работе": PatternFill("solid", fgColor="DBEAFE"),
    "Исполнено": PatternFill("solid", fgColor="D1FAE5"),
    "Просрочено": PatternFill("solid", fgColor="FEE2E2"),
    "Передано": PatternFill("solid", fgColor="E2E8F0"),
}
STATUS_FONTS = {
    "В ожидании": Font(name="Calibri", size=10, bold=True, color="B45309"),
    "В работе": Font(name="Calibri", size=10, bold=True, color="1D4ED8"),
    "Исполнено": Font(name="Calibri", size=10, bold=True, color="047857"),
    "Просрочено": Font(name="Calibri", size=10, bold=True, color="B91C1C"),
    "Передано": Font(name="Calibri", size=10, bold=True, color="475569"),
}

STAGE_FILLS = {
    "early": PatternFill("solid", fgColor="E2E8F0"),
    "start": PatternFill("solid", fgColor="DBEAFE"),
    "critical": PatternFill("solid", fgColor="FEF3C7"),
    "overdue": PatternFill("solid", fgColor="FEE2E2"),
    "waiting": PatternFill("solid", fgColor="EDE9FE"),
    "done": PatternFill("solid", fgColor="D1FAE5"),
}
STAGE_FONTS = {
    "early": Font(name="Calibri", size=10, bold=True, color="475569"),
    "start": Font(name="Calibri", size=10, bold=True, color="1D4ED8"),
    "critical": Font(name="Calibri", size=10, bold=True, color="B45309"),
    "overdue": Font(name="Calibri", size=10, bold=True, color="B91C1C"),
    "waiting": Font(name="Calibri", size=10, bold=True, color="6D28D9"),
    "done": Font(name="Calibri", size=10, bold=True, color="047857"),
}
STAGE_LABELS = {
    "early": "Ещё рано",
    "start": "Пора начинать",
    "critical": "Критично",
    "overdue": "Просрочено",
    "waiting": "Жду ответа",
    "done": "Исполнено",
}

DEFAULT_WIDTHS = {
    "№": 5,
    "Входящий номер": 16,
    "Поток": 12,
    "Дата входящего": 13,
    "Дата поступления": 14,
    "Отправитель": 28,
    "Содержание": 42,
    "Исполнитель": 22,
    "Срок исполнения": 13,
    "Стадия срока": 14,
    "Статус": 13,
    "Тип запроса": 28,
    "Примечание к запросу": 24,
    "Примечание": 24,
    "Отправлено": 16,
    "Кому": 16,
    "Контроль": 14,
    "Отчет": 18,
    "Дата отчета": 12,
}


def _fmt_date(d: Optional[date]) -> str:
    if not d:
        return ""
    return d.strftime("%d.%m.%Y")


def _executor_name(corr: models.Correspondence) -> str:
    if corr.executor:
        return corr.executor.full_name
    if corr.created_by:
        return corr.created_by.full_name
    return "Не назначен"


def _request_type_label(corr: models.Correspondence) -> str:
    if corr.request_type == models.RequestType.OWN_PURPOSES:
        return "Для собственных целей"
    if corr.request_type == models.RequestType.MINISTRY_OF_EDUCATION:
        return "Для Министерства просвещения России"
    return ""


def _stage_for(corr: models.Correspondence, today: Optional[date] = None) -> str:
    today = today or date.today()
    return deadline_stage(
        corr.deadline,
        today=today,
        status=corr.status,
        waiting_external=bool(getattr(corr, "waiting_external", False)),
    )


def _apply_column_widths(ws: Worksheet, headers: Sequence[str]) -> None:
    for idx, header in enumerate(headers, 1):
        width = DEFAULT_WIDTHS.get(header, 16)
        ws.column_dimensions[get_column_letter(idx)].width = width


def _style_status_cell(cell, status: str) -> None:
    fill = STATUS_FILLS.get(status)
    font = STATUS_FONTS.get(status)
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    else:
        cell.font = CELL_FONT
    cell.alignment = CENTER


def _style_stage_cell(cell, stage: str) -> None:
    fill = STAGE_FILLS.get(stage)
    font = STAGE_FONTS.get(stage)
    if fill:
        cell.fill = fill
    if font:
        cell.font = font
    else:
        cell.font = CELL_FONT
    cell.alignment = CENTER
    cell.value = STAGE_LABELS.get(stage, stage)


def _write_title_block(
    ws: Worksheet,
    title: str,
    subtitle: str,
    col_count: int,
) -> int:
    """Пишет заголовок отчёта. Возвращает номер строки заголовков таблицы."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_count)
    sub_cell = ws.cell(row=2, column=1, value=subtitle)
    sub_cell.font = SUBTITLE_FONT
    sub_cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18
    return 4  # пустая строка 3, заголовки на 4


def build_correspondence_workbook(
    correspondences: Iterable[models.Correspondence],
    *,
    sheet_title: str = "Корреспонденция",
    report_title: str = "Реестр входящей корреспонденции",
    include_flow: bool = True,
    include_request: bool = True,
    extra_subtitle: Optional[str] = None,
) -> BytesIO:
    """
    Собирает оформленный xlsx.
    Колонки статуса и стадии срока подсвечены цветом.
    """
    items = list(correspondences)
    today = date.today()
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    headers: List[str] = ["№", "Входящий номер"]
    if include_flow:
        headers.append("Поток")
    headers.extend([
        "Дата входящего",
        "Дата поступления",
        "Отправитель",
        "Содержание",
        "Исполнитель",
        "Срок исполнения",
        "Стадия срока",
        "Статус",
    ])
    if include_request:
        headers.extend(["Тип запроса", "Примечание к запросу"])
    headers.extend(["Отправлено", "Кому", "Контроль", "Отчет", "Дата отчета"])

    subtitle_parts = [
        f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}",
        f"Записей: {len(items)}",
    ]
    if extra_subtitle:
        subtitle_parts.insert(0, extra_subtitle)
    header_row = _write_title_block(ws, report_title, " · ".join(subtitle_parts), len(headers))

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[header_row].height = 22

    status_col = headers.index("Статус") + 1
    stage_col = headers.index("Стадия срока") + 1
    deadline_col = headers.index("Срок исполнения") + 1

    for row_offset, corr in enumerate(items):
        row_idx = header_row + 1 + row_offset
        stage = _stage_for(corr, today)
        status_value = corr.status.value if corr.status else ""

        # flow helper — avoid circular import of extract from router
        flow_val = corr.flow or ""
        if not flow_val and corr.incoming_number and "/" in corr.incoming_number:
            flow_val = corr.incoming_number.split("/", 1)[1].strip() or "без потока"
        elif not flow_val:
            flow_val = "без потока"

        row_data: List[Any] = [row_offset + 1, corr.incoming_number or ""]
        if include_flow:
            row_data.append(flow_val)
        row_data.extend([
            _fmt_date(corr.incoming_date),
            _fmt_date(corr.received_date),
            corr.sender or "",
            corr.content or "",
            _executor_name(corr),
            _fmt_date(corr.deadline),
            STAGE_LABELS.get(stage, stage),
            status_value,
        ])
        if include_request:
            note = corr.request_note or ""
            row_data.extend([_request_type_label(corr), note])
        row_data.extend([
            corr.sent_info or "",
            corr.to_whom or "",
            corr.control or "",
            corr.report or "",
            _fmt_date(corr.report_date),
        ])

        zebra = row_offset % 2 == 1
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = BORDER
            cell.font = CELL_FONT
            if col_idx in (1, deadline_col, stage_col, status_col):
                cell.alignment = CENTER
            else:
                cell.alignment = LEFT
            if zebra and col_idx not in (status_col, stage_col, deadline_col):
                cell.fill = ZEBRA_FILL

        # Статус
        _style_status_cell(ws.cell(row=row_idx, column=status_col), status_value)

        # Стадия срока (+ подсветка самой даты)
        stage_cell = ws.cell(row=row_idx, column=stage_col)
        _style_stage_cell(stage_cell, stage)

        deadline_cell = ws.cell(row=row_idx, column=deadline_col)
        if stage in STAGE_FILLS:
            deadline_cell.fill = STAGE_FILLS[stage]
            deadline_cell.font = STAGE_FONTS.get(stage, CELL_FONT)
        deadline_cell.alignment = CENTER

        # Чуть выше строки с длинным содержанием
        ws.row_dimensions[row_idx].height = 32

    _apply_column_widths(ws, headers)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(headers))}{header_row + max(len(items), 1)}"
    )
    ws.print_title_rows = f"{header_row}:{header_row}"

    # Легенда на втором листе
    legend = wb.create_sheet("Легенда", 1)
    legend["A1"] = "Легенда статусов и стадий срока"
    legend["A1"].font = TITLE_FONT
    legend.merge_cells("A1:C1")

    legend["A3"] = "Статус"
    legend["B3"] = "Цвет"
    legend["A3"].font = Font(bold=True)
    legend["B3"].font = Font(bold=True)
    for i, (name, fill) in enumerate(STATUS_FILLS.items(), 4):
        legend.cell(row=i, column=1, value=name).font = STATUS_FONTS[name]
        c = legend.cell(row=i, column=2, value="")
        c.fill = fill
        c.border = BORDER

    legend["A10"] = "Стадия срока"
    legend["B10"] = "Цвет"
    legend["A10"].font = Font(bold=True)
    legend["B10"].font = Font(bold=True)
    for i, key in enumerate(("early", "start", "critical", "overdue", "waiting", "done"), 11):
        legend.cell(row=i, column=1, value=STAGE_LABELS[key]).font = STAGE_FONTS[key]
        c = legend.cell(row=i, column=2, value="")
        c.fill = STAGE_FILLS[key]
        c.border = BORDER

    legend.column_dimensions["A"].width = 22
    legend.column_dimensions["B"].width = 14

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output
