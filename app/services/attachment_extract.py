"""Извлечение текста из вложений писем (PDF / DOCX / RTF / TXT)."""

from __future__ import annotations

import io
import re
from typing import Optional


_MAX_CHARS = 80_000


def _clip(text: str) -> str:
    text = (text or "").strip()
    if len(text) > _MAX_CHARS:
        return text[:_MAX_CHARS] + "\n…"
    return text


def extract_text_from_bytes(filename: str, data: bytes, content_type: str = "") -> Optional[str]:
    """Вернуть текст вложения или None, если формат не поддерживается."""
    if not data:
        return None
    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    try:
        if name.endswith(".pdf") or "pdf" in ctype:
            return _from_pdf(data)
        if name.endswith(".docx") or "wordprocessingml" in ctype:
            return _from_docx(data)
        if name.endswith(".rtf") or "rtf" in ctype:
            return _from_rtf(data)
        if name.endswith((".txt", ".csv", ".md", ".log")) or ctype.startswith("text/"):
            return _from_plain(data)
        # Старый .doc — без LibreOffice ненадёжно
        if name.endswith(".doc"):
            return None
    except Exception:
        return None
    return None


def _from_plain(data: bytes) -> str:
    for enc in ("utf-8", "cp1251", "latin-1"):
        try:
            return _clip(data.decode(enc))
        except UnicodeDecodeError:
            continue
    return _clip(data.decode("utf-8", errors="ignore"))


def _from_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages[:40]:
        parts.append(page.extract_text() or "")
    return _clip("\n".join(parts))


def _from_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text and c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return _clip("\n".join(parts))


def _from_rtf(data: bytes) -> str:
    from striprtf.striprtf import rtf_to_text

    raw = None
    for enc in ("utf-8", "cp1251", "latin-1"):
        try:
            raw = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        raw = data.decode("utf-8", errors="ignore")
    text = rtf_to_text(raw)
    # Убрать лишние управляющие куски
    text = re.sub(r"\s+", " ", text)
    return _clip(text)
